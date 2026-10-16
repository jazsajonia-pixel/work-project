from pathlib import Path
from docx import Document
from pptx import Presentation
import re
DAILY_HEADER='RHPU12 / GSCHPT DAILY REPORT\nWHAT: {what}\nWHEN: {when}\nWHERE: {where}\nWHO: {who}\nHOW: {how}\n\nVEHICLE RECORDS (MC/MV)\n'
def validate_batch(batch, vehicles, uploads=None, template_types=None):
    blocking=[]; review=[]
    for field,label in [('incident_date','incident date'),('incident_time','incident time'),('location','incident location')]:
        if not batch.get(field): blocking.append(f'Missing required {label}.')
    if not batch.get('period_start') or not batch.get('period_end'): review.append('Weekly reporting period is incomplete; required for weekly templates.')
    if batch.get('period_start') and batch.get('period_end') and batch['period_start']>batch['period_end']: blocking.append('Reporting period start must not be after its end.')
    if not vehicles: blocking.append('No vehicle records have been added.')
    for i,v in enumerate(vehicles,1):
        if v.get('category') not in {'MC','MV'}: blocking.append(f'Vehicle {i} is missing a valid MC/MV category.')
        if not v.get('recovery_date') or not v.get('recovery_place'): blocking.append(f'Vehicle {i} is missing recovery date or place.')
        if not v.get('plate_or_sticker') and not v.get('chassis_number'): review.append(f'Vehicle {i} has no plate/sticker or chassis number; review identity fields.')
        files=[u for u in (uploads or []) if u.get('vehicle_id')==v.get('id')]
        if not files: review.append(f'Vehicle {i} has no attached photographs.')
        if not any(u.get('kind')=='DOCUMENT' for u in files): review.append(f'Vehicle {i} has no possessor ID/license attachment.')
    plates=[(v.get('plate_or_sticker') or '').strip().upper() for v in vehicles if (v.get('plate_or_sticker') or '').strip()]
    if len(plates)!=len(set(plates)): review.append('Duplicate plate or conduction-sticker numbers detected within this batch.')
    stated=batch.get('operational_info','')
    for category,label in [('MC','motorcycles'),('MV','motor vehicles')]:
        m=re.search(rf'(\d+)\s+{label}',stated,re.I); actual=sum(v.get('category')==category for v in vehicles)
        if m and int(m.group(1))!=actual: review.append(f'Stated {label} total ({m.group(1)}) disagrees with listed records ({actual}).')
    if template_types is not None:
        for kind in ('TXT','DOCX','PPTX'):
            if kind not in template_types: review.append(f'No active {kind} template is available; safe fallback will be used.')
    return {'blocking':blocking,'review':review,'can_generate':not blocking}
def daily_txt(batch,vehicles):
    text=DAILY_HEADER.format(what='Vehicle impoundment / recovery',when=f"{batch['incident_date']} {batch['incident_time']}",where=batch['location'],who=batch.get('operational_info') or 'Not provided',how=batch.get('notes') or 'Manual record entry')
    for i,v in enumerate(vehicles,1): text+=f"{i}. {v.get('category','')} | Make/Type/Model: {v.get('make','')} {v.get('vehicle_type','')} {v.get('model','')} | Plate/Sticker: {v.get('plate_or_sticker','')} | Color: {v.get('color','')} | Recovery: {v.get('recovery_date','')} at {v.get('recovery_place','')} | Unit: {v.get('recovering_unit','')} | Violation: {v.get('violation','')}\n"
    return text
def weekly_docx(path,batch,vehicles,template=None):
    doc=Document(str(template)) if template and template.exists() else Document()
    if not template: doc.add_heading('DRAFT — WEEKLY OPERATIONAL REPORT',0); doc.add_paragraph('SAFE FALLBACK: Original DOCX template was not present in the repository.')
    doc.add_paragraph(f"Batch ID: {batch['batch_id']} | Preparation date: {batch['preparation_date']} | Period: {batch.get('period_start','')} to {batch.get('period_end','')}")
    doc.add_paragraph(f"Incident: {batch['incident_date']} {batch['incident_time']} — {batch['location']}")
    table=doc.add_table(rows=1,cols=6); table.style='Table Grid'
    for cell,label in zip(table.rows[0].cells,['Category','Make / Type / Model','Plate / Sticker','Recovery','Unit','Status']): cell.text=label
    for v in vehicles:
        cells=table.add_row().cells
        for cell,value in zip(cells,[v.get('category',''),f"{v.get('make','')} {v.get('vehicle_type','')} {v.get('model','')}",v.get('plate_or_sticker',''),f"{v.get('recovery_date','')} / {v.get('recovery_place','')}",v.get('recovering_unit',''),v.get('status','')]): cell.text=value
    doc.save(path)
def weekly_pptx(path,batch,vehicles,template=None):
    prs=Presentation(str(template)) if template and template.exists() else Presentation()
    if not template:
        slide=prs.slides.add_slide(prs.slide_layouts[0]); slide.shapes.title.text='DRAFT — WEEKLY OPERATIONAL REPORT'; slide.placeholders[1].text='SAFE FALLBACK: Original PPTX template was not present.'
    slide=prs.slides.add_slide(prs.slide_layouts[5]); slide.shapes.title.text=f"Batch {batch['batch_id']} — Vehicle Recoveries"
    body=slide.shapes.add_textbox(500000,1200000,8500000,5000000).text_frame; body.text=f"Period: {batch.get('period_start','')} to {batch.get('period_end','')}\nLocation: {batch['location']}\n\n"+'\n'.join(f"{i}. {v.get('category','')} — {v.get('make','')} {v.get('model','')} — {v.get('plate_or_sticker','')}" for i,v in enumerate(vehicles,1)); prs.save(path)

def weekly_docx_aggregate(path, period_start, period_end, batches, vehicles, template=None):
    doc=Document(str(template)) if template and template.exists() else Document()
    if not template: doc.add_heading('DRAFT — WEEKLY OPERATIONAL REPORT',0); doc.add_paragraph('SAFE FALLBACK: Original DOCX template was not present in the repository.')
    doc.add_paragraph(f"Reporting period: {period_start} to {period_end} | Batches included: {len(batches)} | Vehicles: {len(vehicles)}")
    by_batch={b['batch_id']:[] for b in batches}
    for vehicle in vehicles: by_batch.setdefault(vehicle['batch_id'],[]).append(vehicle)
    for batch in batches:
        doc.add_heading(f"Batch {batch['batch_id']} — {batch['incident_date']} — {batch['location']}",1)
        doc.add_paragraph(batch.get('title') or batch.get('operational_info') or 'Operation details not provided.')
        table=doc.add_table(rows=1,cols=6); table.style='Table Grid'
        for cell,label in zip(table.rows[0].cells,['Category','Make / Type / Model','Plate / Sticker','Recovery','Unit','Status']): cell.text=label
        for v in by_batch.get(batch['batch_id'],[]):
            cells=table.add_row().cells
            for cell,value in zip(cells,[v.get('category',''),f"{v.get('make','')} {v.get('vehicle_type','')} {v.get('model','')}",v.get('plate_or_sticker',''),f"{v.get('recovery_date','')} / {v.get('recovery_place','')}",v.get('recovering_unit',''),v.get('status','')]): cell.text=value
    doc.save(path)

def weekly_pptx_aggregate(path, period_start, period_end, batches, vehicles, template=None):
    prs=Presentation(str(template)) if template and template.exists() else Presentation()
    if not template:
        slide=prs.slides.add_slide(prs.slide_layouts[0]); slide.shapes.title.text='DRAFT — WEEKLY OPERATIONAL REPORT'; slide.placeholders[1].text='SAFE FALLBACK: Original PPTX template was not present.'
    by_batch={b['batch_id']:[] for b in batches}
    for vehicle in vehicles: by_batch.setdefault(vehicle['batch_id'],[]).append(vehicle)
    for batch in batches:
        slide=prs.slides.add_slide(prs.slide_layouts[5]); slide.shapes.title.text=f"Batch {batch['batch_id']} — Vehicle Recoveries"
        rows=by_batch.get(batch['batch_id'],[])
        body=slide.shapes.add_textbox(500000,1200000,8500000,5000000).text_frame
        body.text=f"Reporting period: {period_start} to {period_end}\nIncident: {batch['incident_date']} {batch['incident_time']} — {batch['location']}\nVehicles: {len(rows)}\n\n"+'\n'.join(f"{i}. {v.get('category','')} — {v.get('make','')} {v.get('vehicle_type','')} {v.get('model','')} — {v.get('plate_or_sticker','')}" for i,v in enumerate(rows,1))
    prs.save(path)
